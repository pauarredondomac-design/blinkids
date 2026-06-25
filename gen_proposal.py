from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Page margins
section = doc.sections[0]
section.page_width   = Inches(8.5)
section.page_height  = Inches(11)
section.left_margin  = Inches(1.0)
section.right_margin = Inches(1.0)
section.top_margin    = Inches(0.9)
section.bottom_margin = Inches(0.9)

# ── Colors
GREEN   = RGBColor(0x2E, 0x7D, 0x32)
DKGREEN = RGBColor(0x1B, 0x5E, 0x20)
AMBER   = RGBColor(0xF5, 0x7F, 0x17)
GRAY    = RGBColor(0x42, 0x42, 0x42)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
BLACK   = RGBColor(0x1A, 0x1A, 0x1A)
BLUE    = RGBColor(0x01, 0x57, 0x9B)
LTGREEN = RGBColor(0xC8, 0xE6, 0xC9)
MDGREEN = RGBColor(0xA5, 0xD6, 0xA7)


# ── Helpers
def set_cell_bg(cell, hex_color):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)


def add_run(para, text, bold=False, italic=False, size=11, color=BLACK, font='Calibri'):
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name  = font
    run.font.size  = Pt(size)
    run.font.color.rgb = color
    return run


def heading1(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.name  = 'Calibri'
    run.font.size  = Pt(16)
    run.font.color.rgb = GREEN
    pPr    = p._p.get_or_add_pPr()
    pBdr   = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'),   'single')
    bottom.set(qn('w:sz'),    '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '2E7D32')
    pBdr.append(bottom)
    pPr.append(pBdr)


def heading2(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after  = Pt(3)
    run = p.add_run(text)
    run.bold = True
    run.font.name  = 'Calibri'
    run.font.size  = Pt(12)
    run.font.color.rgb = DKGREEN


def body(text, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after  = Pt(4)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(text)
    run.font.name  = 'Calibri'
    run.font.size  = Pt(10.5)
    run.font.color.rgb = GRAY
    run.italic = italic


def bullet(prefix, text=''):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after  = Pt(3)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.left_indent  = Cm(0.8)
    if text:
        r1 = p.add_run(prefix + ' ')
        r1.bold = True
        r1.font.name  = 'Calibri'
        r1.font.size  = Pt(10.5)
        r1.font.color.rgb = BLACK
        r2 = p.add_run(text)
        r2.font.name  = 'Calibri'
        r2.font.size  = Pt(10.5)
        r2.font.color.rgb = GRAY
    else:
        r1 = p.add_run(prefix)
        r1.font.name  = 'Calibri'
        r1.font.size  = Pt(10.5)
        r1.font.color.rgb = GRAY


# =============================================================
# HEADER
# =============================================================
tbl = doc.add_table(rows=1, cols=2)
tbl.style = 'Table Grid'
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.columns[0].width = Inches(4.5)
tbl.columns[1].width = Inches(2.0)

left  = tbl.cell(0, 0)
right = tbl.cell(0, 1)
set_cell_bg(left,  '2E7D32')
set_cell_bg(right, '1B5E20')

left.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
lp1 = left.paragraphs[0]
lp1.alignment = WD_ALIGN_PARAGRAPH.LEFT
lp1.paragraph_format.space_before = Pt(8)
lp1.paragraph_format.space_after  = Pt(2)
add_run(lp1, 'FinQuest', bold=True, size=26, color=WHITE)

lp2 = left.add_paragraph()
lp2.alignment = WD_ALIGN_PARAGRAPH.LEFT
lp2.paragraph_format.space_after = Pt(2)
add_run(lp2, 'Videojuego Educativo de Finanzas Personales', size=10, color=LTGREEN, italic=True)

lp3 = left.add_paragraph()
lp3.alignment = WD_ALIGN_PARAGRAPH.LEFT
lp3.paragraph_format.space_after = Pt(8)
add_run(lp3, 'Propuesta comercial  |  Mayo 2026', size=9, color=MDGREEN)

right.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
rp1 = right.paragraphs[0]
rp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
rp1.paragraph_format.space_before = Pt(8)
add_run(rp1, 'PROPUESTA DE DESARROLLO', bold=True, size=10, color=WHITE)

rp2 = right.add_paragraph()
rp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
rp2.paragraph_format.space_after = Pt(8)
add_run(rp2, 'v1.0  |  2026-05-22', size=8, color=MDGREEN)

doc.add_paragraph()

# =============================================================
# 1. RESUMEN EJECUTIVO
# =============================================================
heading1('1.  Resumen Ejecutivo')
body(
    'FinQuest es un videojuego educativo orientado a ninos y jovenes de 8 a 13 anos '
    'que ensenha conceptos de finanzas personales de forma ludica: ahorro, presupuesto, '
    'trabajo y mercado, dentro de un mundo de biomas 2D con narrativa RPG. '
    'La plataforma esta construida en Flutter (iOS, Android y Web) con motor de '
    'animaciones Flame y backend Supabase (base de datos, autenticacion, Edge Functions), '
    'disenada exclusivamente para el mercado hispanohablante de Mexico.'
)
body(
    'Este documento detalla el alcance del proyecto, el estado de avance actual '
    'y la propuesta economica para completar el desarrollo completo de la aplicacion.'
)

# =============================================================
# 2. DESCRIPCION DEL PRODUCTO
# =============================================================
heading1('2.  Descripcion del Producto')

heading2('Que es FinQuest?')
body(
    'Una app gamificada donde el nino asume el rol de explorador financiero. '
    'Explora biomas (Bosque, Espacio y mas), ingresa a edificios-modulo para aprender '
    'habilidades financieras: administrar su Bolsa de monedas, responder Preguntas '
    'de educacion financiera, completar Misiones, realizar Trabajos (mini-juegos) '
    'y visitar el Mercado y la Tienda virtual.'
)

heading2('Usuarios objetivo')
bullet('Perfil hijo:', 'Ninjo jugador (8-13 anos) que aprende finanzas mediante juego.')
bullet('Perfil padre/tutor:', 'Supervisa, envia monedas y aprueba compras del nino.')

heading2('Tecnologia empleada')
bullet('Flutter 3.32 / Dart 3 — Frontend multiplataforma (iOS, Android y Web)')
bullet('Flame 1.18 — Motor de animaciones y sprites 2D')
bullet('Supabase — PostgreSQL, autenticacion, Storage y Edge Functions')
bullet('Riverpod 2.5 — Gestion de estado reactivo')
bullet('GoRouter 14 — Navegacion declarativa con proteccion de rutas')
bullet('flutter_animate 4.5 — Animaciones de UI enriquecidas')

# =============================================================
# 3. ALCANCE DEL PROYECTO
# =============================================================
heading1('3.  Alcance del Proyecto')

heading2('Modulos completados (Fase 1 y 2A)')
bullet('Arquitectura base: Flutter, tema Material 3, sistema de colores, tipografia Nunito')
bullet('Base de datos: 15 tablas, RLS activo en todas, migraciones versionadas')
bullet('Autenticacion: Google OAuth + correo/contrasena; flujo splash -> login -> registro')
bullet('Registro completo: seleccion de rol (padre/hijo) -> perfil -> cartera -> personaje')
bullet('Tutorial interactivo: 4 pasos con Juan el Zorro, otorga 200 monedas al completar')
bullet('Mapa Bosque 2D: fondo PNG + 6 edificios posicionados con coordenadas imagen 2048x1143')
bullet('Mi Bolsa (Wallet): distribucion y retiro de monedas entre 5 categorias; slider interactivo')
bullet('Panel Padre: cartera, hijos vinculados y envio de monedas')
bullet('Overlays tutoriales: aparecen la primera vez que se visita cada pantalla')
bullet('Seguridad: RLS estricto; anon key en cliente; compras reales solo por padre via Edge Function')

heading2('Modulos en desarrollo (Fase 2B-2D)')
bullet('HUD de monedas integrado en el mapa Bosque (wallet provider)')
bullet('Pantalla Misiones: listado, progreso y recompensas')
bullet('Pantalla Preguntas: quiz opcion multiple y Verdadero/Falso con animaciones de Juan')
bullet('Trabajos - Mini-juego Vendedor de Frutas: dar cambio correcto con tiempo limite')
bullet('Mercado: listado de articulos virtuales, filtros y compra con monedas')
bullet('Tienda: items cosmeticos para el personaje Juan (sombreros, accesorios)')
bullet('Sistema de evolucion: XP, niveles y cambio de avatar segun progreso')
bullet('Desbloqueo progresivo de edificios segun avance del jugador')

heading2('Fase 3 - Expansion (propuesta futura)')
bullet('Bioma Espacio: mapa ya implementado en CustomPainter, pendiente de contenido')
bullet('Bioma adicional (p. ej. Ciudad Medieval): mecanicas de finanzas avanzadas')
bullet('Notificaciones push: recordatorios de misiones y logros')
bullet('Reportes para padres: analisis del aprendizaje del hijo')
bullet('Modo cooperativo: misiones en equipo entre amigos')

# =============================================================
# 4. PROPUESTA ECONOMICA
# =============================================================
heading1('4.  Propuesta Economica')
body('Precios expresados en Pesos Mexicanos (MXN) con IVA incluido.')

price_data = [
    ('Fase 1 — Base del proyecto (COMPLETADA)',
     'Arquitectura Flutter, BD Supabase, autenticacion, registro, tutorial, seguridad RLS',
     '80 h', '$24,000', True),
    ('Fase 2A — Mapa Bosque + Mi Bolsa (COMPLETADA)',
     'Mapa 2D con assets PNG, cartera con distribucion/retiro, panel padre, overlays tutoriales',
     '40 h', '$12,000', True),
    ('Fase 2B — Misiones + Preguntas (PENDIENTE)',
     'Pantalla misiones con progreso y recompensas; quiz opcion multiple y V/F con animaciones',
     '35 h', '$10,500', False),
    ('Fase 2C — Trabajos + Mercado + Tienda (PENDIENTE)',
     'Mini-juego Vendedor de Frutas; pantalla Mercado con filtros; Tienda de cosmeticos para Juan',
     '45 h', '$13,500', False),
    ('Fase 2D — Evolucion + Desbloqueos (PENDIENTE)',
     'Sistema de XP y niveles, desbloqueo progresivo de edificios, sprites de evolucion',
     '20 h', '$6,000', False),
    ('Fase 3 — Expansion (OPCIONAL)',
     'Bioma adicional, notificaciones push, reportes para padres, modo cooperativo',
     '60 h', '$18,000', False),
]

ptbl = doc.add_table(rows=1, cols=5)
ptbl.style = 'Table Grid'
ptbl.alignment = WD_TABLE_ALIGNMENT.CENTER
widths = [Inches(1.65), Inches(2.60), Inches(0.70), Inches(0.90), Inches(0.65)]
for i, w in enumerate(widths):
    for cell in ptbl.columns[i].cells:
        cell.width = w

# Header
hrow = ptbl.rows[0]
for i, h in enumerate(['Modulo / Fase', 'Descripcion', 'Horas', 'Precio MXN', 'Estado']):
    c = hrow.cells[i]
    set_cell_bg(c, '2E7D32')
    cp = c.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cp.add_run(h)
    run.bold = True
    run.font.name  = 'Calibri'
    run.font.size  = Pt(10)
    run.font.color.rgb = WHITE

for (modulo, desc, horas, precio, done) in price_data:
    row = ptbl.add_row()
    bg = 'F9FBE7' if done else 'FFFFFF'

    c0 = row.cells[0]; set_cell_bg(c0, bg)
    p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r0 = p0.add_run(modulo)
    r0.bold = True; r0.font.name = 'Calibri'; r0.font.size = Pt(9)
    r0.font.color.rgb = DKGREEN if done else BLACK

    c1 = row.cells[1]; set_cell_bg(c1, bg)
    p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r1 = p1.add_run(desc)
    r1.font.name = 'Calibri'; r1.font.size = Pt(9); r1.font.color.rgb = GRAY

    c2 = row.cells[2]; set_cell_bg(c2, bg)
    p2 = c2.paragraphs[0]; p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(horas)
    r2.font.name = 'Calibri'; r2.font.size = Pt(9.5); r2.font.color.rgb = GRAY

    c3 = row.cells[3]; set_cell_bg(c3, bg)
    p3 = c3.paragraphs[0]; p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run(precio)
    r3.bold = True; r3.font.name = 'Calibri'; r3.font.size = Pt(10)
    r3.font.color.rgb = GREEN if done else AMBER

    c4 = row.cells[4]; set_cell_bg(c4, bg)
    p4 = c4.paragraphs[0]; p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status = 'Lista' if done else 'Pendiente'
    r4 = p4.add_run(status)
    r4.font.name = 'Calibri'; r4.font.size = Pt(9)
    r4.font.color.rgb = GREEN if done else AMBER

# Subtotal row
sub_row = ptbl.add_row()
for c in sub_row.cells:
    set_cell_bg(c, 'E8F5E9')
p_sub = sub_row.cells[0].paragraphs[0]
p_sub.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r_sub = p_sub.add_run('Fases 1 a 2D completas (sin Fase 3):')
r_sub.bold = True; r_sub.font.name = 'Calibri'; r_sub.font.size = Pt(9.5); r_sub.font.color.rgb = BLACK
sub_row.cells[1].merge(sub_row.cells[2])
p_sub3 = sub_row.cells[3].paragraphs[0]
p_sub3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_sub3 = p_sub3.add_run('$66,000')
r_sub3.bold = True; r_sub3.font.name = 'Calibri'; r_sub3.font.size = Pt(11); r_sub3.font.color.rgb = GREEN

# Total row
tot_row = ptbl.add_row()
for c in tot_row.cells:
    set_cell_bg(c, '2E7D32')
p_tot = tot_row.cells[0].paragraphs[0]
p_tot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r_tot = p_tot.add_run('TOTAL PROYECTO COMPLETO (con Fase 3):')
r_tot.bold = True; r_tot.font.name = 'Calibri'; r_tot.font.size = Pt(9.5); r_tot.font.color.rgb = WHITE
tot_row.cells[1].merge(tot_row.cells[2])
p_tot3 = tot_row.cells[3].paragraphs[0]
p_tot3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_tot3 = p_tot3.add_run('$84,000')
r_tot3.bold = True; r_tot3.font.name = 'Calibri'; r_tot3.font.size = Pt(12); r_tot3.font.color.rgb = WHITE

doc.add_paragraph()

# ── Modalidades de pago
heading2('Modalidades de Pago')

modal_data = [
    ('Pago por Fases', 'Pago al iniciar y al entregar cada fase por separado.',
     '50% inicio + 50% entrega', 'Sin descuento adicional', 'E3F2FD'),
    ('Pago Total  (Recomendado)', 'Pago del proyecto completo al inicio del desarrollo.',
     '$75,600 MXN', '10% de descuento aplicado', 'E8F5E9'),
    ('Suscripcion Mensual', 'Desarrollo continuo y mantenimiento mes a mes.',
     '$8,500 MXN/mes', 'Minimo 6 meses', 'FFF8E1'),
]

mtbl = doc.add_table(rows=1, cols=3)
mtbl.style = 'Table Grid'
mtbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for c in mtbl.columns:
    c.width = Inches(2.16)

mrow = mtbl.rows[0]
for i, (titulo, sub, precio, nota, bg) in enumerate(modal_data):
    cell = mrow.cells[i]
    set_cell_bg(cell, bg)
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(6)

    r1 = cp.add_run(titulo + '\n')
    r1.bold = True; r1.font.name = 'Calibri'; r1.font.size = Pt(10.5); r1.font.color.rgb = DKGREEN

    r2 = cp.add_run(sub + '\n')
    r2.font.name = 'Calibri'; r2.font.size = Pt(8.5); r2.font.color.rgb = GRAY

    r3 = cp.add_run(precio + '\n')
    r3.bold = True; r3.font.name = 'Calibri'; r3.font.size = Pt(13); r3.font.color.rgb = GREEN

    r4 = cp.add_run(nota)
    r4.italic = True; r4.font.name = 'Calibri'; r4.font.size = Pt(8.5); r4.font.color.rgb = AMBER

    cp.paragraph_format.space_after = Pt(6)

doc.add_paragraph()

# =============================================================
# 5. CRONOGRAMA
# =============================================================
heading1('5.  Cronograma Estimado')

sched = [
    ('Mayo 2026',          'Fases 1 y 2A: Arquitectura, autenticacion, mapa Bosque, Mi Bolsa', 'Completada',    '2E7D32'),
    ('Junio 2026',         'Fase 2B: Misiones y Preguntas', 'En proceso', 'F57F17'),
    ('Jun - Jul 2026',     'Fase 2C: Trabajos, Mercado y Tienda', 'Programado', '01579B'),
    ('Julio 2026',         'Fase 2D: Evolucion de Juan y desbloqueos', 'Programado', '01579B'),
    ('Agosto 2026',        'Fase 3: Expansion (bioma adicional, notificaciones, reportes)', 'Opcional', '4A148C'),
    ('Agosto 2026',        'QA, pruebas en dispositivos fisicos, publicacion App Store y Play Store', 'Entrega final', '1B5E20'),
]

stbl = doc.add_table(rows=1, cols=3)
stbl.style = 'Table Grid'
stbl.alignment = WD_TABLE_ALIGNMENT.CENTER
stbl.columns[0].width = Inches(1.4)
stbl.columns[1].width = Inches(3.7)
stbl.columns[2].width = Inches(1.4)

shdr = stbl.rows[0]
for i, h in enumerate(['Periodo', 'Actividad', 'Estado']):
    set_cell_bg(shdr.cells[i], '2E7D32')
    sp = shdr.cells[i].paragraphs[0]
    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sp.add_run(h)
    sr.bold = True; sr.font.name = 'Calibri'; sr.font.size = Pt(10); sr.font.color.rgb = WHITE

for (periodo, actividad, estado, color_hex) in sched:
    srow = stbl.add_row()
    bg = 'F9FBE7' if 'Completada' in estado else 'FFFFFF'
    for c in srow.cells:
        set_cell_bg(c, bg)

    p0 = srow.cells[0].paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run(periodo)
    r0.font.name = 'Calibri'; r0.font.size = Pt(9.5); r0.font.color.rgb = GRAY

    p1 = srow.cells[1].paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r1 = p1.add_run(actividad)
    r1.font.name = 'Calibri'; r1.font.size = Pt(9.5); r1.font.color.rgb = BLACK

    p2 = srow.cells[2].paragraphs[0]; p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(estado)
    r2.font.name = 'Calibri'; r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(
        int(color_hex[0:2], 16),
        int(color_hex[2:4], 16),
        int(color_hex[4:6], 16),
    )

doc.add_paragraph()

# =============================================================
# 6. ENTREGABLES
# =============================================================
heading1('6.  Entregables por Fase')
bullet('Codigo fuente completo en repositorio privado (GitHub o GitLab)')
bullet('APK firmado para Android + build IPA para iOS')
bullet('Version Web deployada en dominio del cliente (Netlify o Vercel)')
bullet('Base de datos Supabase configurada, migraciones y seed data')
bullet('Panel de administracion basico para el padre')
bullet('Manual de uso en PDF (espanol)')
bullet('Sesion de capacitacion virtual de 1 hora por fase')
bullet('1 mes de soporte tecnico post-entrega por fase incluido')

# =============================================================
# 7. TERMINOS Y CONDICIONES
# =============================================================
heading1('7.  Terminos y Condiciones')
bullet('Precios con IVA. Se acepta transferencia SPEI, deposito bancario o PayPal.')
bullet('El 50% de cada fase se paga antes de iniciar; el 50% restante contra entrega funcional.')
bullet('Se incluyen hasta 2 rondas de revision de diseno por pantalla.')
bullet('Cambios de alcance fuera del acuerdo se cotizan a $450 MXN/hora adicional.')
bullet('El cliente conserva el 100% de los derechos del codigo fuente y los assets.')
bullet('La confidencialidad del proyecto esta cubierta por NDA a solicitud del cliente.')
bullet('Los activos de terceros (fuentes, iconos) son de licencia libre/comercial incluida.')

# =============================================================
# 8. CONTACTO
# =============================================================
heading1('8.  Contacto')
body('Para aceptar esta propuesta o solicitar ajustes comuniquese por cualquiera de los siguientes medios:')
doc.add_paragraph()

ctbl = doc.add_table(rows=1, cols=2)
ctbl.style = 'Table Grid'
ctbl.alignment = WD_TABLE_ALIGNMENT.CENTER
ctbl.columns[0].width = Inches(3.25)
ctbl.columns[1].width = Inches(3.25)

cl = ctbl.cell(0, 0)
set_cell_bg(cl, 'E8F5E9')
cp_l = cl.paragraphs[0]
cp_l.alignment = WD_ALIGN_PARAGRAPH.LEFT
cp_l.paragraph_format.space_before = Pt(6)
add_run(cp_l, 'Desarrollador\n', bold=True, size=11, color=DKGREEN)
add_run(cp_l, 'merinogomez17@gmail.com\n', size=10, color=BLUE)
add_run(cp_l, 'Mexico — disponible por WhatsApp y correo', size=10, color=GRAY)
cp_l.paragraph_format.space_after = Pt(6)

cr = ctbl.cell(0, 1)
set_cell_bg(cr, 'F9FBE7')
cp_r = cr.paragraphs[0]
cp_r.alignment = WD_ALIGN_PARAGRAPH.LEFT
cp_r.paragraph_format.space_before = Pt(6)
add_run(cp_r, 'Vigencia de la propuesta\n', bold=True, size=11, color=DKGREEN)
add_run(cp_r, 'Esta cotizacion es valida por 30 dias\n', size=10, color=GRAY)
add_run(cp_r, 'a partir del 22 de mayo de 2026.', size=10, color=GRAY)
cp_r.paragraph_format.space_after = Pt(6)

doc.add_paragraph()

# ── Footer
fp = doc.add_paragraph()
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.space_before = Pt(16)
rf = fp.add_run('FinQuest  |  Propuesta generada el 22 de mayo de 2026  |  Confidencial')
rf.font.name = 'Calibri'
rf.font.size = Pt(8)
rf.font.color.rgb = RGBColor(0xBD, 0xBD, 0xBD)
rf.italic = True

# ── Save
out = r'C:\apps_IOS_android\Juego economia\finquest\FinQuest_Propuesta_Comercial.docx'
doc.save(out)
print('Guardado en:', out)
